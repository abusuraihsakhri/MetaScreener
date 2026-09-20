import io
import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from config import Config


def export_excel() -> bytes:
    """Generate an Excel workbook and return it as bytes."""
    conn = sqlite3.connect(Config.DB_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4A6CF7", end_color="4A6CF7", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")

    def add_sheet(name, query, params=()):
        cur.execute(query, params)
        cols = [c[0] for c in cur.description]
        ws = wb.create_sheet(name)
        for ci, col in enumerate(cols, 1):
            cell = ws.cell(row=1, column=ci, value=col.replace("_", " ").title())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        for ri, row in enumerate(cur.fetchall(), 2):
            for ci, val in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=val)
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 55)
        ws.freeze_panes = "A2"

    add_sheet("All Papers", "SELECT id,title,authors,year,doi,pmid,screening_decision,screening_confidence FROM papers WHERE is_duplicate=0 ORDER BY id")
    add_sheet("Included", "SELECT id,title,authors,year,doi,pmid FROM papers WHERE is_duplicate=0 AND screening_decision='include' ORDER BY id")
    add_sheet("Excluded", "SELECT id,title,screening_reason FROM papers WHERE is_duplicate=0 AND screening_decision='exclude' ORDER BY id")

    cur.execute("SELECT DISTINCT p.id, p.title, e.field_name, e.field_value FROM papers p JOIN extraction e ON p.id=e.paper_id WHERE p.is_duplicate=0 ORDER BY p.id, e.field_name")
    rows = cur.fetchall()
    if rows:
        ws = wb.create_sheet("Extraction")
        ws.append(["Paper ID", "Title", "Field", "Value"])
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = "A2"

    conn.close()

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def prisma_counts() -> dict:
    conn = sqlite3.connect(Config.DB_PATH)
    cur = conn.cursor()
    identified = cur.execute("SELECT COUNT(*) FROM papers").fetchone()[0]
    duplicates = cur.execute("SELECT COUNT(*) FROM papers WHERE is_duplicate=1").fetchone()[0]
    screened = identified - duplicates
    excluded_screening = cur.execute("SELECT COUNT(*) FROM papers WHERE is_duplicate=0 AND screening_decision='exclude'").fetchone()[0]
    eligible = cur.execute("SELECT COUNT(*) FROM papers WHERE is_duplicate=0 AND screening_decision='include'").fetchone()[0]
    included = cur.execute("SELECT COUNT(*) FROM papers WHERE is_duplicate=0 AND fulltext_decision='include'").fetchone()[0]
    excluded_fulltext = cur.execute("SELECT COUNT(*) FROM papers WHERE is_duplicate=0 AND fulltext_decision='exclude'").fetchone()[0]
    conn.close()
    return {
        "identified": identified,
        "duplicates": duplicates,
        "screened": screened,
        "excluded_screening": excluded_screening,
        "eligible": eligible,
        "included": included,
        "excluded_fulltext": excluded_fulltext,
    }
