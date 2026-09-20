import sqlite3
import json
import logging
import os
import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from config import Config
from modules.ollama_client import Extractor as AIExtractor

DEFAULT_FIELDS = [
    {"name": "study_design",        "label": "Study Design",              "type": "select",  "options": ["Case report", "Case series", "Cohort", "RCT", "Retrospective", "Prospective", "Systematic review", "Other"]},
    {"name": "sample_size",         "label": "Sample Size (n)",           "type": "number"},
    {"name": "age_mean",            "label": "Mean Age (years)",          "type": "number"},
    {"name": "age_range",           "label": "Age Range",                 "type": "text"},
    {"name": "sex_male",            "label": "Male patients (n)",         "type": "number"},
    {"name": "sex_female",          "label": "Female patients (n)",       "type": "number"},
    {"name": "tumor_location",      "label": "Tumor Location / Site",     "type": "text"},
    {"name": "staging",             "label": "Staging / Grading",         "type": "text"},
    {"name": "afp_level",           "label": "AFP Level (ng/mL)",         "type": "text"},
    {"name": "chemotherapy",        "label": "Chemotherapy Regimen",      "type": "text"},
    {"name": "intervention",        "label": "Intervention / Treatment",  "type": "text"},
    {"name": "followup_duration",   "label": "Follow-up Duration",        "type": "text"},
    {"name": "survival_outcome",    "label": "Survival Outcome",          "type": "text"},
    {"name": "response_rate",       "label": "Response Rate / Results",   "type": "text"},
    {"name": "adverse_events",      "label": "Adverse Events",            "type": "text"},
    {"name": "conclusions",         "label": "Author Conclusions",        "type": "textarea"},
    {"name": "notes",               "label": "Reviewer Notes",            "type": "textarea"},
]

class ExtractionManager:
    def __init__(self):
        self.db_path = Config.DB_PATH
        self.ai = AIExtractor()
        self.fields = self._load_fields()

    # Allowed field types and max name length for schema validation (M3)
    _ALLOWED_TYPES = {"text", "number", "select", "textarea"}
    _MAX_FIELD_NAME = 100

    def _validate_fields(self, fields) -> list:
        """Validate extraction field schema — reject malformed entries (M3)."""
        if not isinstance(fields, list):
            logging.warning("extraction_fields.json: root must be a list — using defaults")
            return DEFAULT_FIELDS
        validated = []
        for item in fields:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name", "")).strip()[:self._MAX_FIELD_NAME]
            label = str(item.get("label", name)).strip()[:200]
            ftype = item.get("type", "text")
            if ftype not in self._ALLOWED_TYPES:
                ftype = "text"
            if not name:
                continue
            entry = {"name": name, "label": label, "type": ftype}
            if ftype == "select":
                opts = item.get("options", [])
                entry["options"] = [str(o)[:200] for o in opts if isinstance(o, str)]
            validated.append(entry)
        if not validated:
            logging.warning("extraction_fields.json: no valid fields found — using defaults")
            return DEFAULT_FIELDS
        return validated

    def _load_fields(self):
        fields_path = "extraction_fields.json"
        if os.path.exists(fields_path):
            try:
                with open(fields_path, "r", encoding="utf-8") as f:  # M4: explicit encoding
                    raw = json.load(f)
                return self._validate_fields(raw)  # M3: schema validation
            except (json.JSONDecodeError, OSError) as e:
                logging.error(f"Error loading extraction_fields.json: {e}")
                return DEFAULT_FIELDS
        else:
            try:
                with open(fields_path, "w", encoding="utf-8") as f:
                    json.dump(DEFAULT_FIELDS, f, indent=2)
            except OSError as e:
                logging.error(f"Error writing default extraction_fields.json: {e}")
            return DEFAULT_FIELDS

    def save_fields(self, fields):
        self.fields = fields
        try:
            with open("extraction_fields.json", "w", encoding="utf-8") as f:
                json.dump(fields, f, indent=2)
        except OSError as e:
            logging.error(f"Error saving extraction_fields.json: {e}")

    def get_papers_for_extraction(self):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        try:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT * FROM papers
                WHERE screening_decision = 'include' OR fulltext_decision = 'include'
            """)
            papers = [dict(row) for row in cursor.fetchall()]
            for p in papers:
                cursor.execute("SELECT * FROM extraction WHERE paper_id = ?", (p["id"],))
                rows = cursor.fetchall()
                p["extracted"] = {row["field_name"]: dict(row) for row in rows}
        finally:
            conn.close()   # M7: always close
        return papers

    def get_extraction_for_paper(self, paper_id):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        try:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM extraction WHERE paper_id = ?", (paper_id,))
            rows = cursor.fetchall()
            return {row["field_name"]: dict(row) for row in rows}
        finally:
            conn.close()   # M7

    def save_field_value(self, paper_id, field_name, value, ai_suggested=False, human_verified=True):
        conn = sqlite3.connect(self.db_path)
        try:
            cursor = conn.cursor()
            cursor.execute(
                "SELECT id FROM extraction WHERE paper_id = ? AND field_name = ?",
                (paper_id, field_name),
            )
            row = cursor.fetchone()
            if row:
                cursor.execute(
                    "UPDATE extraction SET field_value=?, ai_suggested=?, human_verified=? WHERE id=?",
                    (value, 1 if ai_suggested else 0, 1 if human_verified else 0, row[0]),
                )
            else:
                cursor.execute(
                    "INSERT INTO extraction (paper_id, field_name, field_value, ai_suggested, human_verified) "
                    "VALUES (?, ?, ?, ?, ?)",
                    (paper_id, field_name, value, 1 if ai_suggested else 0, 1 if human_verified else 0),
                )
            conn.commit()
        finally:
            conn.close()   # M7

    def save_all_fields(self, paper_id, field_dict, ai_suggested=False):
        for name, value in field_dict.items():
            # If AI suggested, human_verified is False
            self.save_field_value(paper_id, name, value, ai_suggested=ai_suggested, human_verified=not ai_suggested)

    def ai_extract_paper(self, paper_id):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        try:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM papers WHERE id = ?", (paper_id,))
            paper = cursor.fetchone()
        finally:
            conn.close()   # M7

        if not paper:
            return {}

        text = ""
        pdf_path = paper["pdf_path"]
        # C4: re-validate path — must be an existing .pdf file (not just "exists")
        if (pdf_path
                and os.path.isfile(pdf_path)
                and os.path.splitext(pdf_path)[1].lower() == ".pdf"):
            try:
                doc = fitz.open(pdf_path)
                for page in doc[:10]:
                    text += page.get_text()
                doc.close()
            except Exception as e:
                logging.error(f"Error reading PDF: {e}")
        
        if not text:
            text = f"Title: {paper['title']}\nAbstract: {paper['abstract']}"

        field_names = [f["name"] for f in self.fields]
        extracted = self.ai.extract_fields(text, field_names)
        
        if extracted:
            self.save_all_fields(paper_id, extracted, ai_suggested=True)
            
        return extracted

    def get_extraction_stats(self):
        conn = sqlite3.connect(self.db_path)
        try:
            cursor = conn.cursor()
            cursor.execute(
                "SELECT COUNT(*) FROM papers WHERE screening_decision='include' OR fulltext_decision='include'"
            )
            total = cursor.fetchone()[0]
            cursor.execute(
                "SELECT paper_id, COUNT(*) FROM extraction "
                "WHERE field_value IS NOT NULL AND field_value != '' GROUP BY paper_id"
            )
            counts = cursor.fetchall()
            field_count = len(self.fields)
            fully       = sum(1 for _, cnt in counts if cnt >= field_count)
            partially   = sum(1 for _, cnt in counts if 0 < cnt < field_count)
            not_started = total - (fully + partially)
            cursor.execute(
                "SELECT COUNT(*) FROM extraction WHERE ai_suggested=1 AND human_verified=0"
            )
            unverified = cursor.fetchone()[0]
        finally:
            conn.close()   # M7
        return {
            "total_for_extraction": total,
            "fully_extracted": fully,
            "partially_extracted": partially,
            "not_started": not_started,
            "ai_suggested_unverified": unverified,
        }

    def export_to_excel(self, output_path):
        wb = Workbook()
        
        # Sheet 1: Extraction Data
        ws1 = wb.active
        ws1.title = "Extraction Data"
        
        paper_headers = ["Title", "Authors", "Year", "DOI"]
        field_labels = [f["label"] for f in self.fields]
        headers = paper_headers + field_labels
        
        ws1.append(headers)
        
        # Style headers
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4A6CF7", end_color="4A6CF7", fill_type="solid")
        for cell in ws1[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws1.row_dimensions[1].height = 22
        
        papers = self.get_papers_for_extraction()
        
        # Fill data
        for i, paper in enumerate(papers, start=2):
            row_data = [
                paper.get("title"), 
                paper.get("authors"), 
                paper.get("year"), 
                paper.get("doi")
            ]
            
            ext_data = paper.get("extracted", {})
            for field in self.fields:
                val = ext_data.get(field["name"], {})
                row_data.append(val.get("field_value", ""))
                
            ws1.append(row_data)
            
            # Formatting for cells
            # Light yellow for AI unverified: #FFFDE7
            # Alternating row fill: white and #F9F9FB
            bg_color = "F9F9FB" if i % 2 == 0 else "FFFFFF"
            row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            
            for j, cell in enumerate(ws1[i], start=1):
                cell.fill = row_fill
                # Fields start after paper_headers (index 4)
                if j > len(paper_headers):
                    field_idx = j - len(paper_headers) - 1
                    field_name = self.fields[field_idx]["name"]
                    data = ext_data.get(field_name, {})
                    if data.get("ai_suggested") and not data.get("human_verified"):
                        cell.fill = PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid")
                        if cell.value:
                            cell.comment = Comment("AI \u2014 unverified", "System")
        
        ws1.freeze_panes = "A2"
        
        # Auto-fit columns
        for col in ws1.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))
                except (TypeError, AttributeError):
                    pass
            adjusted_width = min(max_length + 2, 45)
            ws1.column_dimensions[column].width = adjusted_width

        # Sheet 2: Summary
        ws2 = wb.create_sheet("Summary")
        stats = self.get_extraction_stats()
        ws2.append(["Summary Statistics"])
        ws2.append(["Total Papers for Extraction", stats["total_for_extraction"]])
        ws2.append(["Fully Extracted", stats["fully_extracted"]])
        ws2.append(["Partially Extracted", stats["partially_extracted"]])
        ws2.append(["Not Started", stats["not_started"]])
        ws2.append([])
        ws2.append(["Field Name", "Label", "Coverage (%)"])
        
        # Calculate coverage per field
        for field in self.fields:
            name = field["name"]
            filled_count = 0
            for paper in papers:
                if paper.get("extracted", {}).get(name, {}).get("field_value"):
                    filled_count += 1
            coverage = (filled_count / stats["total_for_extraction"] * 100) if stats["total_for_extraction"] > 0 else 0
            ws2.append([name, field["label"], round(coverage, 1)])

        wb.save(output_path)
