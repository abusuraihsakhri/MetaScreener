import sqlite3
import openpyxl
import logging
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from config import Config
from modules.prisma import PRISMAGenerator

class Exporter:
    def __init__(self):
        self.db_path = Config.DB_PATH

    def export_to_excel(self, output_path, include_all=True, include_only=True, include_excluded=True, include_extraction=True, include_prisma=True):
        try:
            wb = openpyxl.Workbook()
            # Remove default sheet
            default_sheet = wb.active
            wb.remove(default_sheet)
            
            conn = sqlite3.connect(self.db_path)
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()
            
            # Helper to format headers
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4A6CF7", end_color="4A6CF7", fill_type="solid")
            center_align = Alignment(horizontal="center", vertical="center")
            
            def add_data_sheet(sheet_name, query, params=()):
                cursor.execute(query, params)
                columns = [col[0] for col in cursor.description]
                
                sheet = wb.create_sheet(sheet_name)
                
                # Write header
                for col_idx, col_name in enumerate(columns, 1):
                    cell = sheet.cell(row=1, column=col_idx, value=col_name.replace("_", " ").title())
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_align
                
                # Write rows
                row_idx = 2
                for row in cursor.fetchall():
                    for col_idx, val in enumerate(row, 1):
                        sheet.cell(row=row_idx, column=col_idx, value=val)
                    row_idx += 1
                
                # Auto-fit
                for col in sheet.columns:
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        try:
                            if cell.value:
                                max_length = max(max_length, len(str(cell.value)))
                        except (TypeError, AttributeError):
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    sheet.column_dimensions[column].width = adjusted_width
                
                # Freeze top row
                sheet.freeze_panes = "A2"
                return sheet

            if include_all:
                add_data_sheet("All Papers", "SELECT title, authors, year, doi, pmid, source_file, screening_decision, screening_reason, screening_confidence FROM papers")
                
            if include_only:
                add_data_sheet("Included Papers", "SELECT title, authors, year, doi, pmid, source_file, screening_decision, screening_reason FROM papers WHERE screening_decision = 'include'")
                
            if include_excluded:
                add_data_sheet("Excluded Papers", "SELECT title, authors, year, doi, pmid, source_file, screening_decision, screening_reason FROM papers WHERE screening_decision = 'exclude'")
                
            if include_extraction:
                # Dynamic extraction table
                try:
                    cursor.execute("SELECT DISTINCT field_name FROM extraction")
                    fields = [r[0] for r in cursor.fetchall()]
                    
                    if fields:
                        cursor.execute("SELECT DISTINCT paper_id FROM extraction")
                        paper_ids = [r[0] for r in cursor.fetchall()]
                        
                        sheet = wb.create_sheet("Extraction Data")
                        headers = ["Paper ID", "Title"] + fields
                        for col_idx, h in enumerate(headers, 1):
                            cell = sheet.cell(row=1, column=col_idx, value=h)
                            cell.font = header_font
                            cell.fill = header_fill
                            
                        for r_idx, pid in enumerate(paper_ids, 2):
                            cursor.execute("SELECT title FROM papers WHERE id = ?", (pid,))
                            p_title_row = cursor.fetchone()
                            p_title = p_title_row[0] if p_title_row else "Unknown"
                            
                            sheet.cell(row=r_idx, column=1, value=pid)
                            sheet.cell(row=r_idx, column=2, value=p_title)
                            
                            for c_idx, field in enumerate(fields, 3):
                                cursor.execute("SELECT field_value FROM extraction WHERE paper_id = ? AND field_name = ?", (pid, field))
                                val_row = cursor.fetchone()
                                if val_row:
                                    sheet.cell(row=r_idx, column=c_idx, value=val_row[0])
                except Exception as ex:
                    logging.warning(f"Could not export extraction data: {ex}")

            if include_prisma:
                PRISMAGenerator().generate_prisma_excel_sheet(wb)
                
            wb.save(output_path)
            conn.close()
            return {"success": True, "path": output_path, "sheets": wb.sheetnames}
            
        except Exception as e:
            logging.error(f"Excel export failed: {e}")
            return {"success": False, "error": str(e)}

    def generate_summary_text(self):
        nums = PRISMAGenerator().get_prisma_numbers()
        report = []
        report.append("METASCREENER — SYSTEMATIC REVIEW SUMMARY REPORT")
        report.append("="*50)
        report.append(f"Total Records Identified: {nums['identified_total']}")
        report.append(f"  - PubMed: {nums['identified_by_db'].get('PubMed', 0)}")
        report.append(f"  - Embase: {nums['identified_by_db'].get('Embase', 0)}")
        report.append(f"  - Scopus: {nums['identified_by_db'].get('Scopus', 0)}")
        report.append(f"  - WoS: {nums['identified_by_db'].get('WoS', 0)}")
        report.append(f"  - Other: {nums['identified_by_db'].get('Other', 0)}")
        report.append("-" * 30)
        report.append(f"Duplicates Removed: {nums['duplicates_removed']}")
        report.append(f"Unique Records Screened: {nums['screened']}")
        report.append("-" * 30)
        report.append(f"Screening Decisions:")
        report.append(f"  - Included: {nums['sought_retrieval']}")
        report.append(f"  - Excluded: {nums['excluded_screening']}")
        report.append(f"  - Not retrieved: {nums['not_retrieved']}")
        report.append("-" * 30)
        report.append(f"Full-text Eligibility:")
        report.append(f"  - Assessed: {nums['assessed_eligibility']}")
        report.append(f"  - Excluded: {nums['excluded_fulltext']}")
        report.append(f"  - Final Included: {nums['included_final']}")
        report.append("=" * 50)
        return "\n".join(report)

    def generate_ai_summary_text(self):
        from modules.ollama_client import OllamaClient
        import json
        client = OllamaClient()
        stats = self.generate_summary_text()
        
        prompt = f"""[INST] You are an expert medical research assistant writing an executive summary of a systematic review.
Here are the final screening numbers in PRISMA format:

{stats}

Write a concise, professional 2-3 paragraph executive summary of this systematic review's screening and full-text eligibility phase. DO NOT invent fake study topics or names; just summarize the numbers logically.
Return ONLY a valid JSON object with exactly one key "executive_summary" mapping to your summary text string.
[/INST]"""
        
        try:
            res_str = client._generate(prompt, 0.4)
            if not res_str:
                return "Error: Could not retrieve summary from AI."
            data = json.loads(res_str)
            return data.get("executive_summary", str(data))
        except Exception as e:
            return f"Error parsing AI response: {e}\nRaw fallback: {res_str if 'res_str' in locals() else 'None'}"
